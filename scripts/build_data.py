from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from statistics import mean
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SITE_ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "anketolog-report-2180102-survey-1052275-sled-uchastnika-VM1VdZ.xlsx"
DATA_JS = SITE_ROOT / "assets" / "data.js"


URL_RE = re.compile(r"https?://[^\s\"<>]+")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-zА-Яа-я]{2,}")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?7|8)?[\s(.-]*\d{3}[\s).-]*\d{3}[\s.-]*\d{2}[\s.-]*\d{2}(?!\d)")
YEAR_RE = re.compile(r"20\d{2}")


POSITIVE_ATTRIBUTION = {
    "Частично связан — был один из нескольких факторов",
    "В значительной степени связан — «Таврида» стала ключевым толчком",
    "Полностью связан — без «Тавриды» этого не произошло бы",
}
STRONG_ATTRIBUTION = {
    "В значительной степени связан — «Таврида» стала ключевым толчком",
    "Полностью связан — без «Тавриды» этого не произошло бы",
}


ATTRIBUTION_META = {
    "Полностью связан — без «Тавриды» этого не произошло бы": ("full", 4),
    "В значительной степени связан — «Таврида» стала ключевым толчком": ("strong", 3),
    "Частично связан — был один из нескольких факторов": ("partial", 2),
    "Скорее не связан — успех определили другие обстоятельства": ("weak", 1),
    "Не связан совсем": ("none", 0),
}


TRAJECTORY_RULES = [
    {
        "id": "project",
        "label": "Автор проекта",
        "description": "Запуск или развитие собственной инициативы после участия.",
        "keywords": [
            "проект",
            "инициатив",
            "запуск",
            "запуст",
            "создал",
            "создала",
            "реализ",
            "студия",
            "центр",
            "лаборатор",
            "продюс",
        ],
    },
    {
        "id": "career",
        "label": "Карьерный переход",
        "description": "Трудоустройство, новая роль, повышение или смена профессиональной сферы.",
        "keywords": [
            "трудоустр",
            "устроил",
            "работ",
            "должност",
            "повыш",
            "карьер",
            "сменил",
            "смена",
            "переш",
            "переход",
            "театр",
            "компания",
            "организац",
        ],
    },
    {
        "id": "visibility",
        "label": "Публичная видимость",
        "description": "Публикации, выставки, выступления, показы, конкурсы и признание.",
        "keywords": [
            "публикац",
            "статья",
            "сми",
            "интервью",
            "youtube",
            "ютуб",
            "яндекс",
            "журнал",
            "канал",
            "издал",
            "книга",
            "рассказ",
            "сборник",
            "выстав",
            "выступ",
            "концерт",
            "спектакл",
            "показ",
            "премьера",
            "фестиваль",
            "сцена",
            "побед",
            "конкурс",
            "лауреат",
            "призер",
            "призёр",
            "диплом",
            "награ",
            "финалист",
        ],
    },
    {
        "id": "grant",
        "label": "Гранты и поддержка",
        "description": "Финансирование, грантовые конкурсы и институциональная поддержка.",
        "keywords": ["грант", "финансир", "субсид", "поддержк", "фонд"],
    },
    {
        "id": "expert",
        "label": "Экспертность",
        "description": "Переход к роли наставника, эксперта, преподавателя или спикера.",
        "keywords": [
            "настав",
            "эксперт",
            "преподав",
            "спикер",
            "лектор",
            "мастер-класс",
            "обуч",
            "ментор",
            "ведущ",
            "организатор",
        ],
    },
    {
        "id": "network",
        "label": "Коллаборации",
        "description": "Профессиональные связи, команды, партнерства и совместные проекты.",
        "keywords": [
            "знакомств",
            "контакт",
            "партнер",
            "партнёр",
            "коллаб",
            "команд",
            "сообще",
            "связ",
            "нетворк",
        ],
    },
    {
        "id": "identity",
        "label": "Самоопределение",
        "description": "Рост уверенности, переосмысление траектории и карьерных целей.",
        "keywords": [
            "уверенн",
            "страх",
            "смел",
            "осознал",
            "осознала",
            "траектор",
            "мотивац",
            "вдохнов",
            "переосмысл",
            "цель",
        ],
    },
]


NO_RESULT_MARKERS = [
    "ничего из перечисленного не появилось",
    "ничего не появилось",
    "нет результата",
    "не появилось",
    "сложно сказать",
    "не могу выделить",
]


MISSING_RULES = [
    {
        "id": "practice",
        "label": "Практика и разборы",
        "keywords": ["практич", "практики", "разбор", "обратн", "ошиб", "текст", "задани"],
    },
    {
        "id": "aftercare",
        "label": "Постпрограммная поддержка",
        "keywords": ["после", "поддержк", "сопровожд", "куратор", "связь", "продолж"],
    },
    {
        "id": "depth",
        "label": "Глубина и уровень программы",
        "keywords": ["глуб", "уров", "лекц", "эксперт", "наполнен", "професс"],
    },
    {
        "id": "time",
        "label": "Больше времени",
        "keywords": ["времен", "дольше", "длитель", "час"],
    },
    {
        "id": "grant-project",
        "label": "Проекты и гранты",
        "keywords": ["грант", "проект", "финанс", "питч", "реализац"],
    },
]


def cell(row: list, col: int) -> str:
    value = row[col - 1] if col - 1 < len(row) else ""
    if value is None:
        return ""
    return str(value).strip()


def normalize(text: str) -> str:
    return text.lower().replace("ё", "е")


def redact(text: str, limit: int = 360) -> str:
    clean = URL_RE.sub("[ссылка]", text)
    clean = EMAIL_RE.sub("[email]", clean)
    clean = PHONE_RE.sub("[телефон]", clean)
    clean = re.sub(r"\s+", " ", clean).strip()
    if len(clean) > limit:
        return clean[: limit - 1].rstrip() + "…"
    return clean


def urls_from(*texts: str) -> list[str]:
    urls: list[str] = []
    for text in texts:
        for raw in URL_RE.findall(text):
            urls.append(raw.rstrip(").,;"))
    return urls


def selected_options(row: list, cols: list[int], options: list[str]) -> list[str]:
    return [options[col - 1] for col in cols if cell(row, col)]


def classify(text: str, time_to_result: str) -> list[str]:
    normalized = normalize(text)
    tags = []
    for rule in TRAJECTORY_RULES:
        if any(keyword.replace("ё", "е") in normalized for keyword in rule["keywords"]):
            tags.append(rule["id"])
    if time_to_result == "Пока не могу выделить отдельный заметный результат" or any(
        marker in normalized for marker in NO_RESULT_MARKERS
    ):
        tags.append("no-clear-result")
    if not tags:
        tags.append("development")
    return list(dict.fromkeys(tags))


def classify_missing(text: str) -> list[str]:
    normalized = normalize(text)
    tags = []
    for rule in MISSING_RULES:
        if any(keyword.replace("ё", "е") in normalized for keyword in rule["keywords"]):
            tags.append(rule["id"])
    return tags


def impact_values(row: list, impact_cols: range) -> list[float | None]:
    values = []
    for col in impact_cols:
        value = cell(row, col)
        if not value or "Затруд" in value:
            values.append(None)
            continue
        try:
            values.append(float(value.replace(",", ".")))
        except ValueError:
            values.append(None)
    return values


def permission_flags(permission: str) -> dict:
    return {
        "caseAllowed": permission != "Нет",
        "nameAllowed": "с указанием моего имени" in permission,
        "materialsAllowed": "прикрепленных материалов" in permission,
        "anonymousOnly": "обезличенном" in permission,
    }


def main() -> None:
    wb = load_workbook(SOURCE, data_only=True)
    ws = wb.active
    rows = [[value if value is not None else "" for value in row] for row in ws.iter_rows(values_only=True)]
    max_col = max(len(row) for row in rows)
    rows = [row + [""] * (max_col - len(row)) for row in rows]

    header_options = [str(value).strip() for value in rows[2]]
    data_rows = rows[4:]

    response_dates = []
    for row in data_rows:
        raw = cell(row, 2)
        try:
            response_dates.append(datetime.strptime(raw, "%Y.%m.%d %H:%M"))
        except ValueError:
            pass

    program_cols = [21, 22, 23]
    detail_cols = [24, 25, 26]
    direction_cols = list(range(27, 49))
    employment_cols = list(range(51, 57))
    factor_cols = list(range(69, 81))
    impact_cols = range(81, 91)
    motivation_cols = list(range(94, 102))
    support_cols = list(range(102, 109))

    records = []
    case_rows = []
    for index, row in enumerate(data_rows, start=1):
        programs = selected_options(row, program_cols, header_options)
        details = [cell(row, col) for col in detail_cols if cell(row, col)]
        directions = selected_options(row, direction_cols, header_options)
        employment = selected_options(row, employment_cols, header_options)
        factors = selected_options(row, factor_cols, header_options)
        motivations = selected_options(row, motivation_cols, header_options)
        supports = selected_options(row, support_cols, header_options)

        years = [int(year) for text in details for year in YEAR_RE.findall(text)]
        latest_year = max(years) if years else None

        attribution = cell(row, 59)
        attribution_group, attribution_score = ATTRIBUTION_META.get(attribution, ("unknown", None))
        time_to_result = cell(row, 62)

        achievement_items = [cell(row, col) for col in [60, 63, 64, 65] if cell(row, col)]
        achievement_parts = achievement_items + ([cell(row, 68)] if cell(row, 68) else [])
        achievement_text = " ".join(achievement_parts)
        skills_text = cell(row, 67)
        future_text = cell(row, 109)
        missing_text = cell(row, 61)
        all_case_text = " ".join([achievement_text, skills_text, future_text])

        impact = impact_values(row, impact_cols)
        numeric_impact = [value for value in impact if value is not None]
        impact_index = round(mean(numeric_impact), 2) if numeric_impact else None

        recommendation = None
        try:
            recommendation = int(cell(row, 92))
        except ValueError:
            pass

        evidence_urls = urls_from(cell(row, 63), cell(row, 64), cell(row, 65), cell(row, 66), cell(row, 68))
        evidence_hosts = Counter(urlparse(url).netloc for url in evidence_urls)
        private_evidence_count = sum(count for host, count in evidence_hosts.items() if host == "ankt.cc")
        external_evidence_count = len(evidence_urls) - private_evidence_count

        permission = cell(row, 111)
        flags = permission_flags(permission)
        interview = cell(row, 110)
        trajectories = classify(all_case_text, time_to_result)
        missing_tags = classify_missing(missing_text)
        record_id = f"T-{index:03d}"
        participant_name = redact(cell(row, 16), 120)
        display_name = participant_name if flags["nameAllowed"] and participant_name else f"Обезличенный кейс {record_id}"

        if "no-clear-result" in trajectories:
            primary_trajectory = "no-clear-result"
        else:
            priority = ["grant", "career", "project", "expert", "visibility", "network", "identity", "development"]
            primary_trajectory = next((tag for tag in priority if tag in trajectories), trajectories[0])

        case_summary = redact(achievement_text or skills_text or future_text, 420)
        if not flags["caseAllowed"]:
            case_summary = ""
            achievement_items = []

        quality_score = 0
        if attribution in STRONG_ATTRIBUTION:
            quality_score += 30
        elif attribution in POSITIVE_ATTRIBUTION:
            quality_score += 18
        if impact_index:
            quality_score += impact_index * 8
        quality_score += min(private_evidence_count + external_evidence_count, 4) * 5
        if interview == "Да":
            quality_score += 8
        if flags["caseAllowed"]:
            quality_score += 8

        record = {
            "id": record_id,
            "displayName": display_name,
            "nameAllowedInInterface": bool(flags["nameAllowed"] and participant_name),
            "responseDate": cell(row, 2),
            "programs": programs,
            "programDetails": [redact(text, 180) for text in details],
            "programYears": sorted(set(years)),
            "latestYear": latest_year,
            "directions": directions,
            "region": cell(row, 49),
            "gender": cell(row, 17),
            "age": cell(row, 18),
            "professionalLevel": cell(row, 50),
            "employment": employment,
            "primaryActivity": redact(cell(row, 57), 140),
            "additionalActivity": redact(cell(row, 58), 140),
            "attribution": attribution,
            "attributionGroup": attribution_group,
            "attributionScore": attribution_score,
            "positiveAttribution": attribution in POSITIVE_ATTRIBUTION,
            "strongAttribution": attribution in STRONG_ATTRIBUTION,
            "timeToResult": time_to_result,
            "impact": impact,
            "impactIndex": impact_index,
            "factors": factors,
            "repeatParticipation": cell(row, 91),
            "recommendation": recommendation,
            "recommendReason": redact(cell(row, 93), 240),
            "motivations": motivations,
            "supportNeeds": supports,
            "futureLink": redact(future_text, 280),
            "caseSummary": case_summary,
            "achievementItems": [redact(text, 460) for text in achievement_items] if flags["caseAllowed"] else [],
            "skillsSummary": redact(skills_text, 240) if flags["caseAllowed"] else "",
            "publicationsSummary": redact(cell(row, 68), 220) if flags["caseAllowed"] else "",
            "missingSummary": redact(missing_text, 220),
            "missingTags": missing_tags,
            "trajectories": trajectories,
            "primaryTrajectory": primary_trajectory,
            "evidenceCount": len(evidence_urls),
            "privateEvidenceCount": private_evidence_count,
            "externalEvidenceCount": external_evidence_count,
            "evidenceHosts": dict(evidence_hosts),
            "interviewConsent": interview,
            "casePermission": permission,
            **flags,
            "qualityScore": round(quality_score, 1),
        }
        records.append(record)

        if flags["caseAllowed"] and (case_summary or record["skillsSummary"] or record["futureLink"]):
            case_rows.append(record)

    impact_labels = [header_options[col - 1] for col in impact_cols]

    data = {
        "meta": {
            "title": "След участника",
            "subtitle": "Влияние программ «Таврида.АРТ» на профессиональную и творческую траекторию",
            "source": SOURCE.name,
            "responses": len(records),
            "responsePeriod": {
                "from": min(response_dates).strftime("%Y-%m-%d") if response_dates else None,
                "to": max(response_dates).strftime("%Y-%m-%d") if response_dates else None,
            },
            "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M"),
        },
        "dictionary": {
            "programs": [header_options[col - 1] for col in program_cols],
            "directions": [header_options[col - 1] for col in direction_cols],
            "employment": [header_options[col - 1] for col in employment_cols],
            "factors": [header_options[col - 1] for col in factor_cols],
            "impactLabels": impact_labels,
            "motivations": [header_options[col - 1] for col in motivation_cols],
            "supportNeeds": [header_options[col - 1] for col in support_cols],
            "trajectories": TRAJECTORY_RULES
            + [
                {
                    "id": "development",
                    "label": "Развитие в процессе",
                    "description": "Есть движение и планы, но конкретный тип результата требует уточнения.",
                },
                {
                    "id": "no-clear-result",
                    "label": "Пока без явного результата",
                    "description": "Участник пока не выделяет конкретный заметный результат.",
                },
            ],
            "missingThemes": MISSING_RULES,
        },
        "records": records,
    }

    DATA_JS.write_text(
        "window.TAVRIDA_DASHBOARD_DATA = "
        + json.dumps(data, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )

    print(f"Wrote {DATA_JS.relative_to(SITE_ROOT)}")
    print(f"Records: {len(records)}; cases: {len(case_rows)}")


if __name__ == "__main__":
    main()
